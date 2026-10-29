"""
The PPT skill's one entry point: material in, deck out.

    generate(work_dir, template, narration=..., title=..., url=...) -> (deck path, token usage or None)
    revise(work_dir, template, comments, narration=..., title=..., url=...) -> (deck path, changes, token usage)
    preview(work_dir, template) -> slide picture names (work_dir/preview/<template>/NN.png)

work_dir keeps everything a deck needs to be regenerated for free: narration.txt (the material), pages/ and pages.json
(the app explored page by page, only when a link was given - once), deck.json (the story), revisions/ (each review
round's comments, the model's report and the story before it). Switching design reuses the story, so it costs no
tokens; rewrite asks Azure for a new story; recapture explores the app again; revise asks Azure to change the story as
the reviewer's comments say (one call), keeping the same fact checks.

The material is the narration and/or the app tour (each screen's text and how it was reached); a link alone is enough.
"""

import datetime
import json
from pathlib import Path

from ppt import preview as previews
from ppt import source
from ppt.capture import capture_app
from ppt.designs import deck_file, design
from ppt.quality import polish
from ppt.writer import StoryError, add_usage, clean_brief, normalize, revise_story, write_story

STORY_VERSION = 2      # deck.json written by an older version of the writer is rewritten, not reused


def story_is_fresh(work_dir: Path, brief: dict | None = None) -> bool:
    """deck.json was written by this writer version, for this brief (when given), after the material it is based on,
    so it can be reused."""
    story = work_dir / "deck.json"
    if not story.exists():
        return False
    try:
        saved = json.loads(story.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return False
    # both sides cleaned: a brief saved before a field existed still matches
    if saved.get("version") != STORY_VERSION or (brief is not None and
                                                 clean_brief(saved.get("brief")) != clean_brief(brief)):
        return False
    inputs = [p for p in (work_dir / "narration.txt", work_dir / "pages.json") if p.exists()]
    return all(story.stat().st_mtime >= p.stat().st_mtime for p in inputs)


def _pages(work_dir: Path) -> list[dict]:
    f = work_dir / "pages.json"
    pages = json.loads(f.read_text(encoding="utf-8")) if f.exists() else []
    return [p for p in pages if (work_dir / "pages" / p["file"]).exists()]


def _tour(pages: list[dict]) -> list[dict]:
    """The app tour the story is written from: each screen's id, how it was reached and what it shows."""
    tour = [{"screen": p["id"], "kind": p.get("kind") or "page", "via": p.get("via") or p["label"],
             "text": p["text"][:1500 if p.get("kind") == "result" else 900]} for p in pages]
    if tour:     # where the app shows its own name: the browser tab's title and the frame around the pages
        tour[0].update(page_title=pages[0].get("title") or "", app_frame=pages[0].get("frame") or "")
    return tour


def _material(pages: list[dict], parts: list[str], tour: list[dict]) -> str:
    """What the facts are checked against: names and claims - the narration and everything the app's screens show
    (figures are checked against the narration alone: the screens' numbers are demo data, shown in the screenshots)."""
    return "\n".join(parts + [f"{t['via']}: {t['text']}" for t in tour] +
                     [f"{t.get('page_title', '')} {t.get('app_frame', '')}" for t in tour[:1]] +
                     [f"{p['label']}: {'; '.join(p['headings'])}" for p in pages])


def _finish(work_dir: Path, template: str, content: dict, pages: list[dict], parts: list[str], brief: dict,
            log=print) -> Path:
    """The checked story saved as deck.json and built in the design: the talk track says where to show each screen in
    the app, the footer where the material came from."""
    via = {p["id"]: p.get("via") for p in pages}
    for sl in content["slides"]:
        if via.get(sl.get("screen") or "") and "In the app:" not in (sl.get("notes") or ""):
            sl["notes"] = "\n\n".join(x for x in (sl.get("notes"), f"In the app: {via[sl['screen']]}.") if x)
    content["source"] = ("product narration and live app" if parts and pages else "live app walkthrough" if pages
                         else "product narration")
    content["version"], content["brief"] = STORY_VERSION, brief
    (work_dir / "deck.json").write_text(json.dumps(content, indent=2, ensure_ascii=False), encoding="utf-8")
    out = work_dir / deck_file(template)
    shots = {p["id"]: work_dir / "pages" / p["file"] for p in pages}
    n = design(template)(content, shots).save(out)
    log(f"Deck ready: {n} slides ({', '.join(s['type'] for s in content['slides'])}) -> {out.name}")
    return out


def generate(work_dir: Path, template: str, narration: str = "", title: str = "", url: str = "",
             rewrite: bool = False, recapture: bool = False, log=print, brief: dict | None = None,
             try_forms: bool = False) -> tuple[Path, dict | None]:
    """Build the deck in the chosen design; returns its path and the story's token usage (None: story reused).
    brief: who the deck is for, how long, what to emphasise (ppt.writer.AUDIENCES / LENGTHS); a story written for
    another brief is rewritten. try_forms: while exploring, also submit forms that are already filled in."""
    brief = clean_brief(brief)
    work_dir.mkdir(parents=True, exist_ok=True)
    if url and (recapture or not (work_dir / "pages.json").exists()):
        log("Exploring the app page by page" + (" - forms already filled in are submitted" if try_forms else
                                                " (read-only)"))
        capture_app(url, work_dir, log, try_forms=try_forms)
    elif url:
        log("Reusing the explored screens")
    pages = _pages(work_dir) if url else []
    parts = source.clean(narration) if narration else []
    tour = _tour(pages)
    if not parts and not tour:
        raise StoryError("nothing to build the deck from - give a narration of at least a few sentences, or an app "
                         "link whose pages could be opened")
    material = _material(pages, parts, tour)
    screen_ids = [p["id"] for p in pages]
    story = work_dir / "deck.json"
    usage = None
    if not rewrite and story_is_fresh(work_dir, brief):
        log("Reusing the deck's story (deck.json) - no tokens spent")
        saved = json.loads(story.read_text(encoding="utf-8"))
        content = normalize(saved, material, screen_ids, title, "detailed" if saved.get("revised") else brief["length"],
                            brief["instructions"] + saved.get("review_notes", ""), figures="\n".join(parts) +
                            saved.get("review_notes", ""))
        content, _ = polish(content, material.split("\n"), log, repair=False, notes_from=parts)
        content.update({k: saved[k] for k in ("revised", "review_notes") if k in saved})
    else:
        focus = f", focus: {brief['focus']}" if brief["focus"] else ""
        log(f"Material: {len(parts)} narration part(s), {len(tour)} app screen(s), {len(material.split()):,} words; "
            f"for {brief['audience']}, {brief['length']} length{focus}")
        raw, usage = write_story(parts, title, tour or None, log, brief)
        content = normalize(raw, material, screen_ids, title, brief["length"], brief["instructions"],
                            figures="\n".join(parts))
        content, fix = polish(content, material.split("\n"), log, repair=True, notes_from=parts)
        usage = add_usage(usage, fix)
        for old in (work_dir / "revisions").glob("*.json") if (work_dir / "revisions").exists() else []:
            old.unlink()      # a new story: the review starts again
    return _finish(work_dir, template, content, pages, parts, brief, log), usage


def revise(work_dir: Path, template: str, comments: list[dict], narration: str = "", title: str = "", url: str = "",
           log=print, brief: dict | None = None) -> tuple[Path, list[dict], dict]:
    """The story changed as the reviewer's comments say ([{"slide": number as shown, or None, "comment": text}]) in one
    Azure call, checked like a new story (facts from the material and the comments only), rebuilt in the design.
    Returns the deck's path, the model's report of each change, and the call's usage."""
    brief = clean_brief(brief)
    story = work_dir / "deck.json"
    if not story.exists():
        raise StoryError("there is no deck to revise yet - generate it first")
    comments = [{"slide": c.get("slide") if isinstance(c.get("slide"), int) else None,
                 "comment": str(c.get("comment") or "").strip()[:600]} for c in comments if str(c.get("comment") or "").strip()]
    if not comments:
        raise StoryError("write at least one comment")
    current = json.loads(story.read_text(encoding="utf-8"))
    built = work_dir / deck_file(template)
    if built.exists():      # what each commented slide shows: designs add slides (an agenda), so numbers alone mislead
        from pptx import Presentation
        shown = [" | ".join(" ".join(sh.text_frame.text.split()) for sh in s.shapes
                            if sh.has_text_frame and sh.text_frame.text.strip()) for s in Presentation(built).slides]
        heads = [" ".join(str(s.get("title") or s.get("text") or "").split()) for s in current["slides"]]
        for c in comments:
            if not c["slide"] or not 1 <= c["slide"] <= len(shown):
                continue
            text = shown[c["slide"] - 1]
            match = [i for i, h in enumerate(heads) if h and h in text]
            c["about"] = ("title slide" if c["slide"] == 1 else f"slides[{match[0] + 1}]" if match else
                          "agenda" if c["slide"] == 2 else f"the slide showing: {text[:300]}")
    pages = _pages(work_dir) if url else []
    parts = source.clean(narration) if narration else []
    tour = _tour(pages)
    material = _material(pages, parts, tour)
    notes = current.get("review_notes", "") + "\n" + "\n".join(c["comment"] for c in comments)
    log(f"Revising the deck from {len(comments)} comment(s)")
    raw, changes, usage = revise_story(current, comments, parts, tour or None, title, log)
    # the comments are material too (a figure the reviewer gives may be used); a revised deck may grow to 15 slides
    content = normalize(raw, material + "\n" + notes, [p["id"] for p in pages], title, "detailed",
                        brief["instructions"] + notes, figures="\n".join(parts) + notes)
    content, fix = polish(content, (material + "\n" + notes).split("\n"), log, repair=True, notes_from=parts)
    usage = add_usage(usage, fix)
    content.update(revised=True, review_notes=notes)
    rounds = work_dir / "revisions"
    rounds.mkdir(exist_ok=True)
    n = len(list(rounds.glob("*.json"))) + 1
    (rounds / f"{n:02d}.json").write_text(json.dumps({
        "round": n, "at": datetime.datetime.now().isoformat(timespec="seconds"), "comments": comments,
        "changes": changes, "usage": usage, "before": current}, indent=2, ensure_ascii=False), encoding="utf-8")
    done = sum(1 for c in changes if c["status"] == "done")
    log(f"Review round {n}: {done} of {len(changes) or len(comments)} change(s) made")
    return _finish(work_dir, template, content, pages, parts, brief, log), changes, usage


def preview(work_dir: Path, template: str, log=print) -> list[str]:
    """Pictures of every slide of the built deck, for the review (work_dir/preview/<template>/NN.png)."""
    deck = work_dir / deck_file(template)
    if not deck.exists():
        return []
    names = previews.render(deck, work_dir / "preview" / template)
    log(f"Preview: {len(names)} slide pictures")
    return names


def revisions(work_dir: Path) -> list[dict]:
    """Every review round so far: when, the comments and what the model did with each (without the old story)."""
    rounds = work_dir / "revisions"
    out = []
    for f in sorted(rounds.glob("*.json")) if rounds.exists() else []:
        r = json.loads(f.read_text(encoding="utf-8"))
        out.append({k: r.get(k) for k in ("round", "at", "comments", "changes", "usage")})
    return out
