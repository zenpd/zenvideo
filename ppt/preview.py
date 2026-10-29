"""
A picture of every slide of a built deck, for reviewing it in the studio before it is approved - drawn from the .pptx
itself with Pillow (no PowerPoint or LibreOffice needed): backgrounds, shapes, pictures (screenshots, icons), charts
and wrapped text in the deck's own fonts. Close to what PowerPoint shows; small differences in line breaks are possible.

    render(deck.pptx, out_dir) -> ["01.png", "02.png", ...]
"""

import io
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

PX = 120          # pixels per inch: a 13.33 x 7.5 in slide becomes 1600 x 900
EMU = 914400
FONT_DIR = Path(r"C:\Windows\Fonts")
FONTS = {   # family -> (regular, bold, italic) files; others fall back to Calibri
    "Calibri": ("calibri.ttf", "calibrib.ttf", "calibrii.ttf"), "Arial": ("arial.ttf", "arialbd.ttf", "ariali.ttf"),
    "Cambria": ("cambria.ttc", "cambriab.ttf", "cambriai.ttf"),
    "Bookman Old Style": ("BOOKOS.TTF", "BOOKOSB.TTF", "BOOKOSI.TTF"),
    "Segoe UI": ("segoeui.ttf", "segoeuib.ttf", "segoeuii.ttf"),
    "Segoe UI Symbol": ("seguisym.ttf", "seguisym.ttf", "seguisym.ttf"),
}
_fonts: dict = {}


def _px(v) -> int:
    return int(round((v or 0) / EMU * PX))


def _font(name: str, size_pt: float, bold: bool, italic: bool):
    key = (name, round(size_pt * 2), bold, italic)
    if key not in _fonts:
        files = FONTS.get(name, FONTS["Calibri"])
        f = files[1] if bold else files[2] if italic else files[0]
        path = FONT_DIR / f
        if not path.exists():       # e.g. Bookman not installed: Georgia stands in
            path = FONT_DIR / ("georgiab.ttf" if bold else "georgia.ttf")
        _fonts[key] = ImageFont.truetype(str(path), max(1, int(round(size_pt * PX / 72))))
    return _fonts[key]


def _rgba(hexv: str, alpha: float = 1.0) -> tuple:
    return tuple(int(hexv[i:i + 2], 16) for i in (0, 2, 4)) + (int(alpha * 255),)


def _fill(shape) -> tuple | None:
    sp = shape._element.spPr
    sf = sp.find(qn("a:solidFill")) if sp is not None else None
    if sf is None or not len(sf) or sf[0].get("val") is None:
        return None
    clr = sf[0]
    a = clr.find(qn("a:alpha"))
    return _rgba(clr.get("val"), int(a.get("val")) / 100000 if a is not None else 1)


def _text(img: Image.Image, shape) -> None:
    tf = shape.text_frame
    x0, y0, w, h = _px(shape.left), _px(shape.top), _px(shape.width), _px(shape.height)
    inset_l, inset_t = _px(tf.margin_left), _px(tf.margin_top)
    x0, w = x0 + inset_l, max(1, w - inset_l - _px(tf.margin_right))
    d = ImageDraw.Draw(img)
    lines = []
    for para in tf.paragraphs:
        runs = para.runs
        if not runs:
            continue
        r = runs[0]
        size = r.font.size.pt if r.font.size else 18
        f = _font(r.font.name or "Calibri", size, bool(r.font.bold), bool(r.font.italic))
        spc = int(r.font._rPr.get("spc") or 0) / 100 * PX / 72
        try:
            color = _rgba(str(r.font.color.rgb)) if r.font.color and r.font.color.type else (0, 0, 0, 255)
        except AttributeError:
            color = (0, 0, 0, 255)
        width = lambda s: f.getlength(s) + spc * len(s)   # noqa: E731
        for chunk in "".join(x.text for x in runs).split("\n"):
            line = ""
            for word in chunk.split():
                trial = f"{line} {word}".strip()
                if width(trial) <= w or not line:
                    line = trial
                else:
                    lines.append((line, f, color, para.alignment, size * 1.18 * PX / 72, spc))
                    line = word
            lines.append((line, f, color, para.alignment, size * 1.18 * PX / 72, spc))
    total = sum(ln[4] for ln in lines)
    anchor = str(tf.vertical_anchor)
    y = y0 + inset_t + ((h - total) / 2 if "MIDDLE" in anchor else (h - total - inset_t) if "BOTTOM" in anchor else 0)
    for text, f, color, align, lh, spc in lines:
        tw = f.getlength(text) + spc * len(text)
        x = x0 + ((w - tw) / 2 if align and "CENTER" in str(align) else (w - tw) if align and "RIGHT" in str(align) else 0)
        if spc:
            for ch in text:
                d.text((x, y), ch, font=f, fill=color)
                x += f.getlength(ch) + spc
        else:
            d.text((x, y), text, font=f, fill=color)
        y += lh


def _chart(d: ImageDraw.ImageDraw, sh, box) -> None:
    ch = sh.chart
    plot = ch.plots[0]
    cats, vals = list(plot.categories), list(plot.series[0].values)
    kind = str(ch.chart_type)
    x0, y0, x1, y1 = box
    lab = _font("Calibri", 11, False, False)

    def col(fmt, default="888888"):
        try:
            return _rgba(str(fmt.fill.fore_color.rgb))
        except Exception:  # noqa: BLE001 - no colour set: grey
            return _rgba(default)
    if "PIE" in kind or "DOUGHNUT" in kind:
        tot, a = sum(vals) or 1, -90
        cx1 = x0 + (x1 - x0) * 0.62
        r = min(cx1 - x0, y1 - y0) / 2 - 6
        cxm, cym = x0 + (cx1 - x0) / 2, (y0 + y1) / 2
        for i, v in enumerate(vals):
            span = 360 * v / tot
            d.pieslice([cxm - r, cym - r, cxm + r, cym + r], a, a + span, fill=col(plot.series[0].points[i].format))
            a += span
        if "DOUGHNUT" in kind:
            d.ellipse([cxm - r * 0.58, cym - r * 0.58, cxm + r * 0.58, cym + r * 0.58], fill=(255, 255, 255, 255))
        for i, c in enumerate(cats):
            d.rectangle([cx1 + 8, y0 + 20 + i * 18, cx1 + 20, y0 + 32 + i * 18], fill=col(plot.series[0].points[i].format))
            d.text((cx1 + 26, y0 + 18 + i * 18), f"{c} {vals[i]:g}", font=lab, fill=(80, 80, 80, 255))
        return
    m, n = max(vals) or 1, len(vals)
    fill = col(plot.series[0].format)
    if "BAR" in kind:
        bh = (y1 - y0 - 20) / n
        for i, v in enumerate(vals):
            by = y0 + 10 + i * bh
            d.rectangle([x0 + 110, by + bh * 0.2, x0 + 110 + (x1 - x0 - 170) * v / m, by + bh * 0.8], fill=fill)
            d.text((x0 + 6, by + bh * 0.35), str(cats[i])[:16], font=lab, fill=(110, 110, 110, 255))
            d.text((x0 + 116 + (x1 - x0 - 170) * v / m, by + bh * 0.35), f"{v:g}", font=lab, fill=(60, 60, 60, 255))
        return
    bw = (x1 - x0 - 20) / n
    pts = []
    for i, v in enumerate(vals):
        bx = x0 + 10 + i * bw
        top = y1 - 26 - (y1 - y0 - 56) * v / m
        if "LINE" in kind:
            pts.append((bx + bw / 2, top))
        else:
            d.rectangle([bx + bw * 0.2, top, bx + bw * 0.8, y1 - 26], fill=fill)
        d.text((bx + bw * 0.3, top - 16), f"{v:g}", font=lab, fill=(60, 60, 60, 255))
        d.text((bx + bw * 0.15, y1 - 22), str(cats[i])[:14], font=lab, fill=(110, 110, 110, 255))
    if pts:
        try:
            lc = _rgba(str(plot.series[0].format.line.color.rgb))
        except Exception:  # noqa: BLE001
            lc = fill
        d.line(pts, fill=lc, width=3)


def _shape(img: Image.Image, sh, W: int, H: int) -> Image.Image:
    layer = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(layer)
    box = [_px(sh.left), _px(sh.top), _px(sh.left) + _px(sh.width), _px(sh.top) + _px(sh.height)]
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in sh.shapes:
            img = _shape(img, child, W, H)
        return img
    if getattr(sh, "has_chart", False) and sh.has_chart:
        _chart(d, sh, box)
    elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
        iw, ih = im.size
        crop = (int(sh.crop_left * iw), int(sh.crop_top * ih), int(iw - sh.crop_right * iw), int(ih - sh.crop_bottom * ih))
        im = im.crop(crop).resize((max(1, box[2] - box[0]), max(1, box[3] - box[1])), Image.LANCZOS)
        layer.paste(im, (box[0], box[1]), im)
    elif sh.shape_type == MSO_SHAPE_TYPE.LINE or sh._element.tag.endswith("cxnSp"):
        ln = sh._element.spPr.find(qn("a:ln"))
        sf = ln.find(qn("a:solidFill")) if ln is not None else None
        c = sf[0].get("val") if sf is not None and len(sf) else "000000"
        wd = max(1, _px(int(ln.get("w") or 12700))) if ln is not None else 1
        d.line([_px(sh.begin_x), _px(sh.begin_y), _px(sh.end_x), _px(sh.end_y)], fill=_rgba(c), width=wd)
    elif sh.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX):
        f = _fill(sh)
        geom_el = sh._element.spPr.find(qn("a:prstGeom")) if sh._element.spPr is not None else None
        geom = geom_el.get("prst") if geom_el is not None else "rect"
        if f:
            if geom == "ellipse":
                d.ellipse(box, fill=f)
            elif geom == "roundRect":
                d.rounded_rectangle(box, radius=min(16, (box[3] - box[1]) // 3), fill=f)
            elif geom in ("homePlate", "chevron"):
                x0, y0, x1, y1 = box
                dep = int((y1 - y0) * (sh.adjustments[0] if len(sh.adjustments) else 0.5))
                pts = [(x0, y0), (x1 - dep, y0), (x1, (y0 + y1) // 2), (x1 - dep, y1), (x0, y1)]
                if geom == "chevron":
                    pts.append((x0 + dep, (y0 + y1) // 2))
                d.polygon(pts, fill=f)
            elif geom == "triangle":
                d.polygon([(box[0], box[3]), ((box[0] + box[2]) // 2, box[1]), (box[2], box[3])], fill=f)
            else:
                d.rectangle(box, fill=f)
    img = Image.alpha_composite(img, layer)
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip():
        _text(img, sh)
    return img


def _slide(slide, W: int, H: int) -> Image.Image:
    img = Image.new("RGBA", (W, H), (255, 255, 255, 255))
    try:
        img.paste(_rgba(str(slide.background.fill.fore_color.rgb)), (0, 0, W, H))
    except Exception:  # noqa: BLE001 - no solid background: white
        pass
    for sh in slide.shapes:
        img = _shape(img, sh, W, H)
    return img.convert("RGB")


def render(pptx: Path, out_dir: Path) -> list[str]:
    """Every slide as out_dir/NN.png (1600 px wide) and a small thumbnail out_dir/NN.thumb.jpg; returns the PNG names."""
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("*.*"):
        old.unlink()
    prs = Presentation(pptx)
    W, H = _px(prs.slide_width), _px(prs.slide_height)
    names = []
    for i, s in enumerate(prs.slides, 1):
        im = _slide(s, W, H)
        name = f"{i:02d}.png"
        im.save(out_dir / name, optimize=True)
        im.resize((W // 4, H // 4), Image.LANCZOS).save(out_dir / f"{i:02d}.thumb.jpg", quality=86)
        names.append(name)
    return names
