"""Zen Studio: the original product deck - dark violet title slides, clean white content slides (13.333 x 7.5 in)."""

from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE

from ppt.designs.base import ASSETS, RECT, ROUNDED, WHITE, Design


class Zen(Design):
    id, name = "zen", "Zen Studio"
    description = "The Zen Studio product deck: dark violet title slides, clean white content slides, violet accents."
    tags = ["Light", "Product"]
    sample = Path(__file__).resolve().parents[2] / "Zen Studio - Product Deck 1.pptx"
    W, H = 13.333, 7.5
    colors = dict(dk1="1A1030", lt1="FFFFFF", dk2="1A1030", lt2="F4F2FC", accent1="6C5CE7", accent2="8B7FF0",
                  accent3="D8D2F0", accent4="6B6580", accent5="E6E1F7", accent6="F4F2FC", hlink="6C5CE7",
                  folHlink="8B7FF0")
    INK, MUTED, ACCENT, LAVENDER, PALE = "1A1030", "6B6580", "6C5CE7", "8B7FF0", "D8D2F0"
    SURFACE, LINE, RADIUS = "F4F2FC", "E6E1F7", 0.14

    def dark(self):
        s = self.slide(None)
        s.shapes.add_picture(str(ASSETS / "deck_bg.jpg"), 0, 0, self.prs.slide_width, self.prs.slide_height)
        return s

    def logo(self, s, x, y):
        self.shape(s, ROUNDED, x, y, 0.62, 0.62, self.ACCENT)
        tri = self.shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, x + 0.22, y + 0.17, 0.21, 0.19, WHITE)
        tri.rotation = 90

    def frame(self, kind, kicker, title):
        s = self.slide()
        self.text(s, self.c["app_name"], 0.5, 7.12, 4.0, 0.3, 9, self.MUTED)
        self.text(s, str(self.n), 12.4, 7.12, 0.5, 0.3, 9, self.MUTED, align="right")
        if kicker:
            self.text(s, kicker.upper(), 0.5, 0.5, 9.0, 0.32, 12.5, self.ACCENT, bold=True, tracking=1)
        if not title:
            return s, (0.5, 1.15, 12.33, 5.7)
        size = self.fit(title, 12.33, 1.05, 30, 22, bold=True)
        self.text(s, title, 0.5, 0.95, 12.33, 1.05, size, self.INK, bold=True)
        top = self.bottom(title, 0.95, 12.33, size, bold=True) + 0.4
        return s, (0.5, top, 12.33, 6.85 - top)

    def cover(self):
        c = self.c
        s = self.dark()
        self.logo(s, 6.36, 1.55)
        self.text(s, c["app_name"], 0.5, 2.35, 12.33, 1.1, 54, WHITE, bold=True, align="center", floor=36)
        self.text(s, c["tagline"], 0.5, 3.5, 12.33, 0.6, 19, self.LAVENDER, align="center", floor=15)
        self.text(s, c["pitch"], 2.17, 4.35, 9.0, 0.8, 13.5, self.PALE, align="center", floor=12)

    def closing(self, sl):
        s = self.dark()
        self.logo(s, 6.36, 1.4)
        size = self.fit(sl["title"], 12.33, 1.6, 40, 28, bold=True)
        self.text(s, sl["title"], 0.5, 2.25, 12.33, 1.6, size, WHITE, bold=True, align="center")
        end = self.bottom(sl["title"], 2.25, 12.33, size, bold=True)
        self.text(s, sl["subtitle"], 1.5, end + 0.25, 10.33, 0.6, 16, self.LAVENDER, align="center", floor=12)
        steps = sl["next_steps"]
        if steps:
            cols, cw = self.columns(1.5, 10.33, len(steps), 0.3)
            for i, ((cx, _), st) in enumerate(zip(cols, steps)):
                self.shape(s, ROUNDED, cx, 5.15, cw, 0.75, "2A1E5C", radius=0.2)
                self.text(s, f"{i + 1}  {st}", cx + 0.2, 5.15, cw - 0.4, 0.75, 13, WHITE, anchor="middle",
                          align="center", floor=10)
