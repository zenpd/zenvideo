"""
Slide drawing shared by every design: primitives in inches, and the deck's slide types laid out generically from a
design's tokens (colours, fonts, card style, marker style). A design (one module each in ppt/designs/) supplies its
tokens, its page frame (background, header, footer - frame() returns the content box), its cover and closing, and may
override any slide type with a layout of its own (e.g. Consulting's chevron process).

Slide types (see ppt/writer.py): statement, stats, problem, before_after, process, feature_grid, feature_focus, quote,
benefits, closing - plus an agenda the builder adds to longer decks (no tokens).
"""

import datetime
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from ppt.fonts import fit, lines
from ppt.icons import icon_png

ASSETS = Path(__file__).resolve().parent.parent / "assets"
SAMPLES = ASSETS / "samples"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
RECT, ROUNDED, OVAL = MSO_SHAPE.RECTANGLE, MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.OVAL
WHITE = "FFFFFF"
QUOTE = "“"


def shadow(color: str, alpha: float, blur: float = 4, dist: float = 1) -> str:
    return (f'<a:effectLst xmlns:a="{A_NS}"><a:outerShdw blurRad="{int(blur * 12700)}" dist="{int(dist * 12700)}" '
            f'dir="5400000" algn="bl" rotWithShape="0"><a:srgbClr val="{color}"><a:alpha val="{int(alpha * 100000)}"/>'
            f'</a:srgbClr></a:outerShdw></a:effectLst>')


class Design:
    id = name = description = ""
    tags: list[str] = []
    sample: Path | None = None
    W, H = 10.0, 5.625
    HEAD = BODY = "Calibri"
    HEAD_BOLD = True
    colors: dict[str, str] = {}     # theme colour scheme written into the file: dk1 lt1 dk2 lt2 accent1-6 hlink folHlink
    # tokens used by the generic slide layouts
    INK, MUTED, BG = "1F2937", "6B7280", WHITE
    ACCENT = "2F6BFF"                # numbers, strips, quote marks
    ACCENTS: list[str] = []          # per-item colours (empty: ACCENT for all)
    STRONG = None                    # filled emphasis: markers, the "after" column, callouts (None: ACCENT)
    ON_STRONG = WHITE
    SURFACE: str | None = "F3F4F6"   # card fill (None: no cards, text on the page)
    LINE = "E5E7EB"
    RADIUS: float | None = 0.1       # card corner radius (None: square)
    CARD_SHADOW: str | None = None
    MARKER = "circle"                # "circle": icon in a filled circle; "number": a big number in the accent colour

    def __init__(self, content: dict, shots: dict[str, Path] | None = None):
        self.c, self.shots = content, shots or {}
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(self.W), Inches(self.H)
        self.blank = self.prs.slide_layouts[6]
        self.k = self.W / 10        # type scale: generic sizes are set for a 10 in wide page
        self.n = 0
        self.focus = 0               # feature_focus slides so far (they alternate sides)
        self.date = datetime.date.today().strftime("%B %Y")
        self._theme()

    # ── file ──
    def _theme(self) -> None:
        part = self.prs.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        ns = {"a": A_NS}
        scheme = root.find(".//a:clrScheme", ns)
        scheme.set("name", self.name)
        for key, value in self.colors.items():
            el = scheme.find(f"a:{key}", ns)
            if el is not None:
                el[:] = []
                etree.SubElement(el, qn("a:srgbClr")).set("val", value)
        for tag, face in (("majorFont", self.HEAD), ("minorFont", self.BODY)):
            root.find(f".//a:fontScheme/a:{tag}/a:latin", ns).set("typeface", face)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def save(self, out: Path) -> int:
        self.build()
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.core_properties.title = self.c["app_name"]
        self.prs.save(out)
        return self.n

    def build(self) -> None:
        slides = [s for s in self.c["slides"] if s["type"] != "closing"]
        closing = next(s for s in self.c["slides"] if s["type"] == "closing")
        self.cover()
        self.note(self.c.get("cover_notes"))
        sections = list(dict.fromkeys(s["kicker"] for s in slides if s.get("kicker")))
        custom = (self.c.get("brief") or {}).get("audience") == "custom"   # "slide 1" must stay the first slide
        if len(slides) >= 7 and len(sections) >= 3 and not custom:
            self.agenda(sections[:6])
            self.note("Today: " + "; ".join(sections[:6]) + ".")
        for s in slides:
            getattr(self, s["type"])(s)
            self.note(s.get("notes"))
        self.closing(closing)
        self.note(closing.get("notes"))

    def wrap_up(self, sl) -> tuple[str, list[str]]:
        """What a closing list shows: the next steps when the material gives some, else the deck's key messages."""
        if sl["next_steps"]:
            return "Next steps", sl["next_steps"]
        keys = [s.get("title") or s.get("text") for s in self.c["slides"]
                if s["type"] in ("statement", "feature_focus", "before_after", "benefits")]
        return "Summary", [k for k in keys if k][:3]

    def note(self, text: str | None) -> None:
        """Speaker notes for the slide just added."""
        if text:
            self.prs.slides[-1].notes_slide.notes_text_frame.text = text

    # ── what a design provides ──
    def frame(self, kind: str, kicker: str, title: str | None):
        """A content slide in this design's chrome; returns (slide, (x, y, w, h) content box below the title)."""
        raise NotImplementedError

    def cover(self):
        raise NotImplementedError

    def closing(self, sl):
        raise NotImplementedError

    # ── primitives ──
    @property
    def strong(self) -> str:
        return self.STRONG or self.ACCENT

    def accent(self, i: int) -> str:
        return self.ACCENTS[i % len(self.ACCENTS)] if self.ACCENTS else self.ACCENT

    def slide(self, bg: str | None = WHITE):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        if bg:
            s.background.fill.solid()
            s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
        return s

    def fit(self, text, w, h, pt, floor, bold=False, font=None, tracking=0) -> float:
        return fit(text, w, h, pt, floor, bold, family=font or self.BODY, tracking=tracking)

    def fit_all(self, texts, w, h, pt, floor, bold=False, font=None) -> float:
        """One size for a row of labels, so they don't come out in mixed sizes."""
        return min([self.fit(t, w, h, pt, floor, bold, font) for t in texts if t] or [pt])

    def bottom(self, text, y, w, pt, bold=False, font=None) -> float:
        """y where text set at y, pt, width w ends."""
        return y + lines(text, pt, bold, w, font or self.BODY) * pt * 1.2 / 72

    def text(self, s, text, x, y, w, h, pt, color, bold=False, italic=False, font=None, align="left", anchor="top",
             floor=None, tracking=0, spacing=1.0):
        """Text box with no inner margins; floor shrinks the size (down to floor) until the text fits w x h."""
        font = font or self.BODY
        if not text:
            return None
        if floor is not None:
            pt = self.fit(text, w, h, pt, floor, bold, font, tracking)
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(max(w, 0.1)), Inches(max(h, 0.1)))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        for k, line in enumerate(str(text).split("\n")):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            para.line_spacing = spacing
            run = para.add_run()
            run.text = line
            run.font.name, run.font.size, run.font.bold, run.font.italic = font, Pt(pt), bold, italic
            run.font.color.rgb = RGBColor.from_string(color)
            if tracking:
                run.font._rPr.set("spc", str(int(tracking * 100)))
        return box

    def head(self, s, text, x, y, w, h, pt, color=None, floor=None, align="left", anchor="top"):
        """Text in the design's heading font (bold or not, per the design)."""
        return self.text(s, text, x, y, w, h, pt, color or self.INK, bold=self.HEAD_BOLD, font=self.HEAD,
                         floor=floor, align=align, anchor=anchor)

    def shape(self, s, kind, x, y, w, h, fill, line=None, line_w=1.0, radius=None, alpha=None, shadow=None):
        """radius: corner radius in inches (rounded rectangles); alpha: fill opacity 0-1; shadow: effectLst xml."""
        shp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
        if alpha is not None:
            clr = shp._element.spPr.find(qn("a:solidFill"))[0]
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))
        if line:
            shp.line.color.rgb = RGBColor.from_string(line)
            shp.line.width = Pt(line_w)
        else:
            shp.line.fill.background()
        if radius is not None and kind == ROUNDED:
            shp.adjustments[0] = min(0.5, radius / min(w, h))
        if shadow:
            shp._element.spPr.append(parse_xml(shadow))
        else:
            shp.shadow.inherit = False
        return shp

    def rule(self, s, x1, y1, x2, y2, color, width=0.75):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        ln.line.color.rgb = RGBColor.from_string(color)
        ln.line.width = Pt(width)
        return ln

    def badge(self, s, label, x, y, d, fill, color=WHITE, pt=13, kind=OVAL, font=None, radius=None):
        """A filled circle (or square) with a centred label."""
        self.shape(s, kind, x, y, d, d, fill, radius=radius)
        self.text(s, label, x, y, d, d, pt, color, bold=True, align="center", anchor="middle", font=font)

    def card(self, s, x, y, w, h, fill=None, strip=None):
        """A content card in the design's style (nothing for designs without cards); strip: accent bar on top."""
        fill = fill or self.SURFACE
        if fill:
            self.shape(s, ROUNDED if self.RADIUS else RECT, x, y, w, h, fill, radius=self.RADIUS,
                       shadow=self.CARD_SHADOW if fill == self.SURFACE else None)
        if strip:
            self.shape(s, RECT, x + (self.RADIUS or 0), y, w - 2 * (self.RADIUS or 0), 0.06 * self.k, strip)

    def icon(self, s, name, x, y, d, color) -> bool:
        png = icon_png(name, color) if name else None
        if png:
            s.shapes.add_picture(str(png), Inches(x), Inches(y), Inches(d), Inches(d))
        return bool(png)

    def marker(self, s, x, y, d, icon="", label="", color=None):
        """An item's marker: its icon in a filled circle (the number when there is no icon), or - for MARKER
        "number" designs - the number large in the accent colour."""
        if self.MARKER == "number":
            self.head(s, label, x, y - d * 0.1, d * 2.2, d * 1.1, d * 72 * 0.72, color or self.ACCENT)
            return
        fill = color or self.strong
        self.shape(s, OVAL, x, y, d, d, fill)
        if not self.icon(s, icon, x + d * 0.24, y + d * 0.24, d * 0.52, self.ON_STRONG):
            self.text(s, label, x, y, d, d, d * 72 * 0.36, self.ON_STRONG, bold=True, align="center", anchor="middle")

    def picture(self, s, path: Path, x, y, w, h, border=None, radius=None):
        """Screenshot cropped to the box (sides trimmed evenly, or the top of the page kept)."""
        pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        iw, ih = Image.open(path).size
        want = w / h
        if iw / ih > want:
            pic.crop_left = pic.crop_right = (1 - want * ih / iw) / 2
        else:
            pic.crop_bottom = 1 - (iw / want) / ih
        if radius:
            geom = pic._element.spPr.find(qn("a:prstGeom"))
            geom.set("prst", "roundRect")
            geom[:] = []
            gd = etree.SubElement(etree.SubElement(geom, qn("a:avLst")), qn("a:gd"))
            gd.set("name", "adj")
            gd.set("fmla", f"val {int(min(0.5, radius / min(w, h)) * 100000)}")
        if border:
            pic.line.color.rgb = RGBColor.from_string(border)
            pic.line.width = Pt(0.75)
        return pic

    def narrow(self, w) -> bool:
        """A content area too narrow for three cards side by side (e.g. beside a navigation panel)."""
        return w < 7.5 * self.k

    def item_rows(self, s, x, y, w, h, items):
        """Items as stacked rows - marker, title, body - for narrow content areas."""
        k = self.k
        rh = h / len(items)
        tsize = self.fit_all([it["title"] for it in items], w - 0.8 * k, 0.4 * k, 15 * k, 11 * k, self.HEAD_BOLD,
                             self.HEAD)
        for i, it in enumerate(items):
            ry = y + i * rh
            if i:
                self.rule(s, x, ry - 0.05 * k, x + w, ry - 0.05 * k, self.LINE)
            self.marker(s, x, ry + 0.08 * k, 0.55 * k, it.get("icon", ""), f"{i + 1:02d}", self.accent(i))
            self.head(s, it["title"], x + 0.8 * k, ry + 0.04 * k, w - 0.8 * k, 0.4 * k, tsize)
            self.text(s, it["body"], x + 0.8 * k, ry + 0.45 * k, w - 0.8 * k, rh - 0.55 * k, 12.5 * k, self.MUTED,
                      floor=9.5 * k)

    def columns(self, x, w, n, gap):
        cw = (w - gap * (n - 1)) / n
        return [(x + i * (cw + gap), cw) for i in range(n)], cw

    # ── slide types ──
    def statement(self, sl):
        k = self.k
        s, (x, y, w, h) = self.frame("statement", sl["kicker"], None)
        tw, ty = w * 0.88 - 0.35 * k, y + h * 0.1
        size = self.fit(sl["text"], tw, h * 0.55, 40 * k, 24 * k, self.HEAD_BOLD, self.HEAD)
        end = self.bottom(sl["text"], ty, tw, size, self.HEAD_BOLD, self.HEAD)
        self.shape(s, RECT, x, ty + 0.06 * k, 0.09 * k, max(0.3, end - ty - 0.06 * k), self.ACCENT)
        self.head(s, sl["text"], x + 0.35 * k, ty, tw, h * 0.55, size)
        self.text(s, sl["support"], x + 0.35 * k, end + 0.3 * k, tw * 0.92, y + h - end - 0.3 * k, 16 * k, self.MUTED,
                  floor=11 * k)

    def stats(self, sl):
        k = self.k
        s, (x, y, w, h) = self.frame("stats", sl["kicker"], sl["title"])
        cols, cw = self.columns(x, w, len(sl["stats"]), 0.3 * k)
        ch = min(h - 0.1 * k, 2.5 * k)
        cy = y + (h - ch) / 2
        vsize = self.fit_all([st["value"] for st in sl["stats"]], cw - 0.6 * k, 1.0 * k, 50 * k, 26 * k, True, self.HEAD)
        for i, ((cx, _), st) in enumerate(zip(cols, sl["stats"])):
            self.card(s, cx, cy, cw, ch, strip=self.accent(i))
            self.text(s, st["value"], cx + 0.3 * k, cy + 0.35 * k, cw - 0.6 * k, 1.0 * k, vsize, self.accent(i),
                      bold=True, font=self.HEAD)
            self.text(s, st["label"], cx + 0.3 * k, cy + 1.45 * k, cw - 0.6 * k, ch - 1.6 * k, 14 * k, self.INK,
                      floor=10 * k)

    def problem(self, sl):
        k = self.k
        s, (x, y, w, h) = self.frame("problem", sl["kicker"], sl["title"])
        top = y
        if sl["intro"]:
            self.text(s, sl["intro"], x, y, w, 0.55 * k, 14 * k, self.MUTED, floor=11 * k)
            top += 0.72 * k
        if self.narrow(w):
            return self.item_rows(s, x, top, w, y + h - top, sl["points"])
        cols, cw = self.columns(x, w, len(sl["points"]), 0.3 * k)
        ch = y + h - top
        tsize = self.fit_all([p["title"] for p in sl["points"]], cw - 0.6 * k, 0.62 * k, 16 * k, 11 * k,
                             self.HEAD_BOLD, self.HEAD)
        for i, ((cx, _), p) in enumerate(zip(cols, sl["points"])):
            self.card(s, cx, top, cw, ch)
            self.marker(s, cx + 0.3 * k, top + 0.3 * k, 0.6 * k, p["icon"], f"{i + 1:02d}")
            self.head(s, p["title"], cx + 0.3 * k, top + 1.08 * k, cw - 0.6 * k, 0.62 * k, tsize)
            self.text(s, p["body"], cx + 0.3 * k, top + 1.72 * k, cw - 0.6 * k, ch - 1.9 * k, 12.5 * k, self.MUTED,
                      floor=9.5 * k)

    def before_after(self, sl):
        k = self.k
        s, (x, y, w, h) = self.frame("before_after", sl["kicker"], sl["title"])
        gap = 0.7 * k
        cw = (w - gap) / 2
        ax = x + cw + gap
        hh = 0.55 * k
        self.card(s, x, y, cw, h)
        self.card(s, ax, y, cw, h, fill=self.strong)
        self.text(s, sl["before_label"].upper(), x + 0.3 * k, y + 0.18 * k, cw - 0.6 * k, 0.32 * k, 11 * k, self.MUTED,
                  bold=True, tracking=1.5, floor=8)
        self.text(s, sl["after_label"].upper(), ax + 0.3 * k, y + 0.18 * k, cw - 0.6 * k, 0.32 * k, 11 * k,
                  self.ON_STRONG, bold=True, tracking=1.5, floor=8)
        rows = sl["rows"]
        rh = (h - hh - 0.15 * k) / len(rows)
        size = self.fit_all([r["before"] for r in rows] + [r["after"] for r in rows], cw - 1.0 * k, rh - 0.15 * k,
                            14 * k, 10 * k)
        for i, r in enumerate(rows):
            ry = y + hh + i * rh
            self.rule(s, x + 0.3 * k, ry, x + cw - 0.3 * k, ry, self.LINE)
            self.rule(s, ax + 0.3 * k, ry, ax + cw - 0.3 * k, ry, self.ON_STRONG, 0.5)
            self.text(s, "–", x + 0.3 * k, ry, 0.3 * k, rh, size, self.MUTED, anchor="middle")
            self.text(s, r["before"], x + 0.65 * k, ry, cw - 0.95 * k, rh, size, self.MUTED, anchor="middle")
            if not self.icon(s, "check", ax + 0.3 * k, ry + rh / 2 - 0.13 * k, 0.26 * k, self.ON_STRONG):
                self.text(s, "✓", ax + 0.3 * k, ry, 0.3 * k, rh, size, self.ON_STRONG, anchor="middle")
            self.text(s, r["after"], ax + 0.65 * k, ry, cw - 0.95 * k, rh, size, self.ON_STRONG, bold=True,
                      anchor="middle")
        d = 0.5 * k
        self.badge(s, "→", x + cw + (gap - d) / 2, y + h / 2 - d / 2, d, self.ACCENT, self.BG, pt=16 * k)

    def process(self, sl):
        k = self.k
        s, (x, y, w, h) = self.frame("process", sl["kicker"], sl["title"])
        top = y
        if sl["intro"]:
            self.text(s, sl["intro"], x, y, w, 0.5 * k, 14 * k, self.MUTED, floor=11 * k)
            top += 0.75 * k
        steps = sl["steps"]
        col = w / len(steps)
        d = 0.7 * k
        cy = top + 0.1 * k
        self.rule(s, x + col / 2, cy + d / 2, x + col * (len(steps) - 0.5), cy + d / 2, self.LINE, 2)
        nsize = self.fit_all([st["name"] for st in steps], col - 0.2 * k, 0.55 * k, 15 * k, 10 * k, True)
        for i, st in enumerate(steps):
            cx = x + i * col
            self.badge(s, str(i + 1), cx + (col - d) / 2, cy, d, self.accent(i) if self.ACCENTS else self.strong,
                       self.ON_STRONG, pt=17 * k)
            self.text(s, st["name"], cx + 0.1 * k, cy + d + 0.25 * k, col - 0.2 * k, 0.55 * k, nsize, self.INK,
                      bold=True, align="center")
            self.text(s, st["body"], cx + 0.12 * k, cy + d + 0.85 * k, col - 0.24 * k, y + h - (cy + d + 0.85 * k),
                      12 * k, self.MUTED, align="center", floor=9 * k)

    def feature_grid(self, sl):
        k = self.k
        s, (x, y, w, h) = self.frame("feature_grid", sl["kicker"], sl["title"])
        items = sl["items"]
        if self.narrow(w) and len(items) == 3:
            return self.item_rows(s, x, y, w, h, items)
        ncols = 2 if self.narrow(w) or len(items) == 4 else 3
        nrows = (len(items) + ncols - 1) // ncols
        gap = 0.25 * k
        cols, cw = self.columns(x, w, ncols, gap)
        ch = (h - gap * (nrows - 1)) / nrows
        if nrows == 1:
            tsize = self.fit_all([it["title"] for it in items], cw - 0.6 * k, 0.62 * k, 16 * k, 11 * k,
                                 self.HEAD_BOLD, self.HEAD)
        else:
            tsize = self.fit_all([it["title"] for it in items], cw - 1.1 * k, 0.5 * k, 14.5 * k, 10 * k,
                                 self.HEAD_BOLD, self.HEAD)
        for i, it in enumerate(items):
            (cx, _), cy = cols[i % ncols], y + (i // ncols) * (ch + gap)
            self.card(s, cx, cy, cw, ch)
            if nrows == 1:
                self.marker(s, cx + 0.3 * k, cy + 0.3 * k, 0.6 * k, it["icon"], f"{i + 1:02d}", self.accent(i))
                self.head(s, it["title"], cx + 0.3 * k, cy + 1.1 * k, cw - 0.6 * k, 0.62 * k, tsize)
                self.text(s, it["body"], cx + 0.3 * k, cy + 1.75 * k, cw - 0.6 * k, ch - 1.95 * k, 12.5 * k,
                          self.MUTED, floor=9.5 * k)
            elif ch < 1.3 * k:   # short cards (6 items in a narrow area): icon left, title and body beside it
                self.marker(s, cx + 0.15 * k, cy + 0.15 * k, 0.42 * k, it["icon"], f"{i + 1:02d}", self.accent(i))
                self.head(s, it["title"], cx + 0.7 * k, cy + 0.1 * k, cw - 0.85 * k, 0.34 * k, 13.5 * k,
                          floor=10 * k)
                self.text(s, it["body"], cx + 0.7 * k, cy + 0.46 * k, cw - 0.85 * k, ch - 0.54 * k, 11 * k,
                          self.MUTED, floor=8.5 * k)
            else:   # icon and title on one line, the body underneath at full card width
                self.marker(s, cx + 0.22 * k, cy + 0.22 * k, 0.5 * k, it["icon"], f"{i + 1:02d}", self.accent(i))
                self.head(s, it["title"], cx + 0.88 * k, cy + 0.22 * k, cw - 1.1 * k, 0.5 * k, tsize, anchor="middle")
                self.text(s, it["body"], cx + 0.22 * k, cy + 0.85 * k, cw - 0.44 * k, ch - 0.98 * k, 11.5 * k,
                          self.MUTED, floor=9 * k)

    def feature_focus(self, sl):
        k = self.k
        s, (x, y, w, h) = self.frame("feature_focus", sl["kicker"], sl["title"])
        self.focus += 1
        shot = self.shots.get(sl.get("screen") or "")
        shot = shot if shot and shot.exists() else None
        tw, pw = (w * 0.38, w * 0.58) if shot else (w * 0.55, w * 0.4)    # a screenshot gets the room to be read
        tx, px = (x, x + w - pw) if self.focus % 2 else (x + w - tw, x)
        top = y
        if sl["body"]:
            size = self.fit(sl["body"], tw, h * 0.45, 15 * k, 11 * k)
            self.text(s, sl["body"], tx, top, tw, h * 0.45, size, self.INK)
            top = self.bottom(sl["body"], top, tw, size) + 0.3 * k
        if sl["points"]:
            ph = (y + h - top) / len(sl["points"])
            psize = self.fit_all(sl["points"], tw - 0.5 * k, ph - 0.1 * k, 13.5 * k, 10 * k)
            for i, p in enumerate(sl["points"]):
                py = top + i * ph
                if not self.icon(s, "check", tx, py + 0.02 * k, 0.26 * k, self.ACCENT):
                    self.shape(s, OVAL, tx + 0.06 * k, py + 0.08 * k, 0.14 * k, 0.14 * k, self.ACCENT)
                self.text(s, p, tx + 0.45 * k, py, tw - 0.45 * k, ph - 0.1 * k, psize, self.INK)
        if shot:
            ph = min(h, pw * 9 / 16)
            self.picture(s, shot, px, y + (h - ph) / 2, pw, ph, border=self.LINE, radius=self.RADIUS)
            return
        self.callout(s, px, y, pw, h, sl["callout"] or sl["title"])

    def callout(self, s, x, y, w, h, text):
        """The takeaway panel: a filled block with a big quote mark."""
        k = self.k
        self.card(s, x, y, w, h, fill=self.strong)
        pad = 0.35 * k
        self.text(s, QUOTE, x + pad, y + pad * 0.5, w - 2 * pad, 1.05 * k, 60 * k, self.ON_STRONG, bold=True,
                  font=self.HEAD)
        self.head(s, text, x + pad, y + 1.25 * k, w - 2 * pad, h - 1.25 * k - pad, 22 * k, self.ON_STRONG,
                  floor=13 * k)

    def quote(self, sl):
        k = self.k
        s, (x, y, w, h) = self.frame("quote", "", None)
        self.text(s, QUOTE, x, y - 0.2 * k, 1.5 * k, 1.4 * k, 110 * k, self.ACCENT, bold=True, font=self.HEAD)
        self.text(s, sl["text"], x + 0.1 * k, y + 1.05 * k, w * 0.86, h - 1.65 * k, 30 * k, self.INK, italic=True,
                  font=self.HEAD, floor=18 * k)
        if sl["attribution"]:
            self.text(s, f"— {sl['attribution']}", x + 0.1 * k, y + h - 0.45 * k, w, 0.4 * k, 14 * k,
                      self.MUTED, bold=True)

    def benefits(self, sl):
        k = self.k
        s, (x, y, w, h) = self.frame("benefits", sl["kicker"], sl["title"])
        cols, cw = self.columns(x, w, len(sl["items"]), 0.3 * k)
        tsize = self.fit_all([it["title"] for it in sl["items"]], cw - 0.6 * k, 0.62 * k, 17 * k, 12 * k,
                             self.HEAD_BOLD, self.HEAD)
        for i, ((cx, _), it) in enumerate(zip(cols, sl["items"])):
            self.card(s, cx, y, cw, h, strip=self.accent(i))
            if it["metric"]:
                self.text(s, it["metric"], cx + 0.3 * k, y + 0.3 * k, cw - 0.6 * k, 0.85 * k, 40 * k, self.accent(i),
                          bold=True, font=self.HEAD, floor=24 * k)
            else:
                self.marker(s, cx + 0.3 * k, y + 0.35 * k, 0.62 * k, it["icon"], f"{i + 1:02d}", self.accent(i))
            self.head(s, it["title"], cx + 0.3 * k, y + 1.3 * k, cw - 0.6 * k, 0.62 * k, tsize)
            self.text(s, it["body"], cx + 0.3 * k, y + 1.95 * k, cw - 0.6 * k, h - 2.15 * k, 12.5 * k, self.MUTED,
                      floor=9.5 * k)

    def palette(self, n: int) -> list[str]:
        """n distinct colours from the design: its accents, then lighter tints of them."""
        base = self.ACCENTS or [self.ACCENT, self.strong]
        out = list(base)
        mix = 0.35
        while len(out) < n:
            for c in base:
                rgb = [int(c[i:i + 2], 16) for i in (0, 2, 4)]
                out.append("".join(f"{int(v + (255 - v) * mix):02X}" for v in rgb))
            mix += 0.2
        return out[:n]

    def chart(self, sl):
        """A native PowerPoint chart (editable in PowerPoint) in the design's colours, with the takeaway beside it."""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION

        k = self.k
        s, (x, y, w, h) = self.frame("chart", sl["kicker"], sl["title"])
        side = bool(sl["takeaway"]) and not self.narrow(w)
        cw = w * 0.64 if side else w
        ch = h if side or not sl["takeaway"] else h - 0.75 * k
        kind = sl["chart"]
        xl = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
              "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT}[kind]
        data = CategoryChartData()
        data.categories = sl["categories"]
        unit = sl["unit"].replace('"', "")
        fmt = '#,##0.##' + (f'"{unit}"' if unit in ("%",) else f'" {unit}"' if unit else "")
        data.add_series(sl["series"] or sl["title"] or "Value", sl["values"], number_format=fmt)
        self.card(s, x, y, cw, ch)
        pad = 0.15 * k if self.SURFACE else 0
        frame = s.shapes.add_chart(xl, Inches(x + pad), Inches(y + pad), Inches(cw - 2 * pad), Inches(ch - 2 * pad), data)
        c = frame.chart
        c.font.name, c.font.size = self.BODY, Pt(11 * k)
        c.font.color.rgb = RGBColor.from_string(self.INK)
        plot = c.plots[0]
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.number_format, labels.number_format_is_linked = fmt, False
        labels.font.size, labels.font.bold = Pt(12 * k), True
        colors = self.palette(len(sl["values"]))
        if kind in ("pie", "doughnut"):
            c.has_legend = True
            c.legend.position, c.legend.include_in_layout = XL_LEGEND_POSITION.RIGHT, False
            c.legend.font.size = Pt(11 * k)
            labels.font.color.rgb = RGBColor.from_string(WHITE if kind == "pie" else self.INK)
            if kind == "pie":
                labels.position = XL_LABEL_POSITION.CENTER
            for i, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(colors[i])
                point.format.line.color.rgb = RGBColor.from_string(self.SURFACE or self.BG)
            if kind == "doughnut":
                hole = plot._element.find(qn("c:holeSize"))
                if hole is not None:
                    hole.set("val", "58")
        else:
            c.has_legend = False
            labels.font.color.rgb = RGBColor.from_string(self.INK)
            labels.position = XL_LABEL_POSITION.OUTSIDE_END if kind != "line" else XL_LABEL_POSITION.ABOVE
            series = plot.series[0]
            if kind == "line":
                series.format.line.color.rgb = RGBColor.from_string(self.ACCENT)
                series.format.line.width = Pt(3)
                series.smooth = False
            else:
                plot.gap_width = 70
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = RGBColor.from_string(self.ACCENT)
            va, ca = c.value_axis, c.category_axis
            va.visible, va.has_major_gridlines = False, False
            ca.tick_labels.font.size = Pt(11 * k)
            ca.tick_labels.font.color.rgb = RGBColor.from_string(self.MUTED)
            ca.format.line.color.rgb = RGBColor.from_string(self.LINE)
            ca.has_major_gridlines = False
            if kind == "bar":
                ca.reverse_order = True     # horizontal bars read top-down in the given order
        if sl["takeaway"]:
            if side:
                self.callout(s, x + cw + 0.3 * k, y, w - cw - 0.3 * k, h, sl["takeaway"])
            else:
                self.text(s, sl["takeaway"], x, y + ch + 0.15 * k, w, 0.6 * k, 14 * k, self.INK, bold=True,
                          floor=10 * k)

    def agenda(self, sections):
        k = self.k
        s, (x, y, w, h) = self.frame("agenda", "Agenda", "What we'll cover")
        ncols = 2 if len(sections) > 4 else 1
        per = (len(sections) + ncols - 1) // ncols
        cols, cw = self.columns(x, w, ncols, 0.5 * k)
        rh = min(0.85 * k, h / per)
        for i, name in enumerate(sections):
            (cx, _), ry = cols[i // per], y + (i % per) * rh
            self.rule(s, cx, ry, cx + cw, ry, self.LINE)
            self.text(s, f"{i + 1:02d}", cx, ry, 0.8 * k, rh, 24 * k, self.ACCENT, bold=self.HEAD_BOLD,
                      font=self.HEAD, anchor="middle")
            self.text(s, name, cx + 0.9 * k, ry, cw - 0.9 * k, rh, 18 * k, self.INK, anchor="middle", floor=12 * k)
