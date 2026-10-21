"""
Pitch deck: a recorded demo -> deck.pptx, in the style of the Zen Studio product deck.

  story       Azure OpenAI writes the pitch (problem, solution, features, benefits, closing) from the demo's own
              narration and transcript - grounded in them, no invented numbers or customers.
  screenshots real frames of the recording (raw.mp4), taken at the end of the segment each slide talks about.
  layout      the reference deck's geometry, colours and type: dark gradient title/closing slides, white content
              slides with a purple kicker, 30 pt title, grey body, framed screenshot + caption, footer with page number.
              Text is measured with Calibri's real metrics and shrunk (never below the floor) so nothing overflows.

    python -m demo deck <script.yaml> [--build-dir ...]      (the recording must exist: run record first)
"""

import json
import re
import subprocess
import tempfile
from pathlib import Path

from PIL import Image, ImageFont

from demo import media
from demo.script import Demo

ASSETS = Path(__file__).parent / "assets"
BACKGROUND = ASSETS / "deck_bg.jpg"
FONT_DIR = Path(r"C:\Windows\Fonts")

# Reference deck palette and type (Zen Studio - Product Deck).
INK, MUTED, ACCENT, LAVENDER, PALE, TINT, WHITE = "1A1030", "6B6580", "6C5CE7", "8B7FF0", "D8D2F0", "F4F2FC", "FFFFFF"
FRAME = "E6E1F7"
# Reference screenshot frame: soft shadow under the white frame.
SHADOW_XML = ('<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:outerShdw blurRad="152400" '
              'dist="50800" dir="5400000" algn="bl" rotWithShape="0"><a:srgbClr val="1A1030"><a:alpha val="30000"/>'
              '</a:srgbClr></a:outerShdw></a:effectLst>')
FONT = "Calibri"
SLIDE_W, SLIDE_H = 13.333, 7.5


class DeckError(RuntimeError):
    pass


# ── content (Azure OpenAI) ───────────────────────────────────────────────────

PROMPT = """You write the slide content for a short product pitch deck about the web application shown in a narrated
demo video. You get the demo's segments (id, what is shown, narration) and the full transcript.

Return a JSON object:
{
 "app_name": product name, max 4 words,
 "tagline": max 9 words,
 "pitch": one sentence, max 28 words,
 "problem": {"title": max 6 words, "intro": max 30 words,
             "points": [3 x {"title": max 5 words, "body": max 22 words}],
             "screenshot": segment id, "caption": max 12 words},
 "solution": {"title": max 9 words, "body": max 40 words, "flow": [3-6 stage names, max 3 words each],
              "screenshot": segment id},
 "features": [3-6 x {"kicker": max 3 words, "title": max 8 words, "body": max 40 words,
                     "screenshot": segment id, "caption": max 12 words}],
 "benefits": {"title": max 6 words, "cards": [3 x {"title": max 4 words, "body": max 20 words}]},
 "closing": {"title": max 7 words, "subtitle": max 12 words}
}

Rules:
- Use ONLY facts stated in the narration/transcript. Never invent numbers, customers, prices or integrations.
- Features follow the order of the demo; each is a capability the demo actually shows.
- "screenshot" must be the id of the segment whose screen best shows that slide's point (ids are listed below);
  use each id at most once where you can.
- Plain, confident product language. No exclamation marks. Titles without a final period.
- Return only the JSON object."""


def _segments_for_prompt(demo: Demo, lang: str) -> list[dict]:
    out = []
    for s in demo.segments:
        shown = [a.describe() for a in s.actions if a.kind in ("click", "select", "fill", "scroll_to", "upload")]
        out.append({"id": s.id, "shows": shown[:6], "narration": s.say.get(lang, "")})
    return out


def write_content(demo: Demo, lang: str, transcript: str, log=print) -> dict:
    from demo.draft import azure_config

    cfg = azure_config()
    if not cfg:
        raise DeckError("Azure OpenAI is not configured (AZURE_OPENAI_* in .env) - it writes the deck's story.")
    from openai import AzureOpenAI

    client = AzureOpenAI(azure_endpoint=cfg["endpoint"], api_key=cfg["api_key"], api_version=cfg["api_version"])
    payload = {"title": demo.title, "url": demo.url, "segments": _segments_for_prompt(demo, lang),
               "transcript": transcript[:12000]}
    log(f"Writing the deck's story with Azure OpenAI ({cfg['deployment']})...")
    resp = client.chat.completions.create(
        model=cfg["deployment"], temperature=0.3, response_format={"type": "json_object"},
        messages=[{"role": "system", "content": PROMPT},
                  {"role": "user", "content": json.dumps(payload, ensure_ascii=False)}])
    try:
        content = json.loads(resp.choices[0].message.content or "{}")
    except json.JSONDecodeError as e:
        raise DeckError(f"Azure OpenAI returned invalid JSON: {e}") from e
    return normalize(content, demo)


def normalize(c: dict, demo: Demo) -> dict:
    """Fill gaps and clamp counts so the builder never sees a missing key, and point every screenshot at a real
    segment (falling back to segments in demo order)."""
    ids = [s.id for s in demo.segments]
    narrated = [s.id for s in demo.segments if s.say] or ids
    used: list[str] = []

    def shot(want: str | None) -> str:
        if want in ids and want not in used:
            used.append(want)
            return want
        free = [i for i in narrated if i not in used] or narrated
        pick = free[min(len(free) - 1, len(used) % max(1, len(free)))]
        used.append(pick)
        return pick

    def text(v, default=""):
        return re.sub(r"\s+", " ", str(v or default)).strip()

    c = dict(c or {})
    c["app_name"] = text(c.get("app_name"), demo.title)
    c["tagline"] = text(c.get("tagline"))
    c["pitch"] = text(c.get("pitch"))
    p = dict(c.get("problem") or {})
    p["points"] = [{"title": text(x.get("title")), "body": text(x.get("body"))}
                   for x in (p.get("points") or []) if isinstance(x, dict)][:3]
    p.update(title=text(p.get("title"), "The problem"), intro=text(p.get("intro")),
             caption=text(p.get("caption")), screenshot=shot(p.get("screenshot")))
    c["problem"] = p
    s = dict(c.get("solution") or {})
    s.update(title=text(s.get("title"), "The solution"), body=text(s.get("body")),
             flow=[text(x) for x in (s.get("flow") or []) if text(x)][:6], screenshot=shot(s.get("screenshot")))
    c["solution"] = s
    feats = [f for f in (c.get("features") or []) if isinstance(f, dict)][:6]
    c["features"] = [{"kicker": text(f.get("kicker")), "title": text(f.get("title")), "body": text(f.get("body")),
                      "caption": text(f.get("caption")), "screenshot": shot(f.get("screenshot"))} for f in feats]
    b = dict(c.get("benefits") or {})
    b["cards"] = [{"title": text(x.get("title")), "body": text(x.get("body"))}
                  for x in (b.get("cards") or []) if isinstance(x, dict)][:3]
    b["title"] = text(b.get("title"), "Why it matters")
    c["benefits"] = b
    cl = dict(c.get("closing") or {})
    c["closing"] = {"title": text(cl.get("title"), c["app_name"]), "subtitle": text(cl.get("subtitle"), c["tagline"])}
    return c


# ── screenshots from the recording ───────────────────────────────────────────

STILL = 0.4          # mean abs pixel change (96x54 grey) below which the screen counts as settled
SAME = 2.0           # mean abs difference (64x36 grey) below which two screenshots are the same picture
NEIGHBOURS = 2       # a replacement screenshot comes from a segment at most this far away in the demo (stays on topic)


def _settled_time(motion, fps: int, start: float, end: float) -> float:
    """Latest moment in [start, end] whose last 0.5 s were still - the screen after a step has played out, not a
    spinner or a half-loaded reply. Falls back to just before the end."""
    import numpy as np
    a, b = int(start * fps), min(int(end * fps), len(motion) - 1)
    win = max(1, int(0.5 * fps))
    for k in range(b, a + win, -1):
        if np.max(motion[k - win:k]) < STILL:
            return k / fps
    return max(start + 0.5, end - 0.6)


def screenshots(demo: Demo, lang: str, out_dir: Path, log=print) -> dict[str, Path]:
    """One settled frame of raw.mp4 per segment."""
    from demo import retime
    rec_dir = demo.lang_dir(lang)
    raw, timeline = rec_dir / "raw.mp4", rec_dir / "timeline.json"
    if not raw.exists() or not timeline.exists():
        raise DeckError(f"no recording in {rec_dir} - record the video first")
    tl = json.loads(timeline.read_text(encoding="utf-8"))
    motion = retime.motion(raw, rec_dir / "motion")
    by_index = {s.index: s.id for s in demo.segments}
    shots = {}
    for seg in tl["segments"]:
        sid = by_index.get(seg["index"])
        if sid is None:
            continue
        t = _settled_time(motion, tl["fps"], seg["start"], seg["end"])
        path = out_dir / f"shot-{seg['index']:03d}.png"
        media.run_ffmpeg(["-ss", f"{t:.3f}", "-i", raw, "-frames:v", "1", path])
        shots[sid] = path
    log(f"Took {len(shots)} screenshot(s) from the recording")
    return shots


def dedupe(content: dict, shots: dict[str, Path], order: list[str], log=print) -> None:
    """No two slides show the same picture (two segments that never left the same screen): a duplicate is swapped
    for the screenshot of a nearby segment (so it stays on topic) that isn't a duplicate itself; if there is none,
    the original stays."""
    import numpy as np

    thumbs = {sid: np.asarray(Image.open(p).convert("L").resize((64, 36)), dtype=np.float32)
              for sid, p in shots.items()}

    def same(a, b):
        return float(np.abs(thumbs[a] - thumbs[b]).mean()) < SAME

    slots = [content["problem"], content["solution"], *content["features"]]
    used: list[str] = []
    for n, slot in enumerate(slots):
        sid = slot["screenshot"]
        wanted_later = {x["screenshot"] for x in slots[n + 1:]}
        if sid in thumbs and any(same(sid, u) for u in used) and sid in order:
            k = order.index(sid)
            near = sorted((c for c in order[max(0, k - NEIGHBOURS):k + NEIGHBOURS + 1] if c in thumbs),
                          key=lambda c: abs(order.index(c) - k))
            alt = next((c for c in near if c not in used and c not in wanted_later
                        and not any(same(c, u) for u in used)), None)
            if alt:
                log(f"  '{slot.get('title', '')[:40]}': same picture as an earlier slide - using {alt} instead")
                sid = slot["screenshot"] = alt
        used.append(sid)


# ── text fitting ─────────────────────────────────────────────────────────────

_fonts: dict[tuple[bool, float], ImageFont.FreeTypeFont] = {}


def _font(bold: bool, pt: float) -> ImageFont.FreeTypeFont:
    key = (bold, pt)
    if key not in _fonts:
        name = "calibrib.ttf" if bold else "calibri.ttf"
        try:
            _fonts[key] = ImageFont.truetype(str(FONT_DIR / name), int(round(pt * 10)))   # 10 px per pt
        except OSError:
            _fonts[key] = ImageFont.load_default(int(round(pt * 10)))
    return _fonts[key]


def _lines(text: str, pt: float, bold: bool, width_in: float) -> int:
    font, max_px, total = _font(bold, pt), width_in * 72 * 10, 0
    for para in text.split("\n"):
        line, n = "", 1
        for word in para.split():
            trial = f"{line} {word}".strip()
            if font.getlength(trial) <= max_px or not line:
                line = trial
            else:
                n, line = n + 1, word
        total += n
    return total


def fit(text: str, w: float, h: float, pt: float, floor: float, bold: bool = False, spacing: float = 1.2) -> float:
    """Largest size from pt down to floor (0.5 pt steps) at which text wraps inside w x h inches."""
    size = pt
    while size > floor and _lines(text, size, bold, w) * size * spacing / 72 > h:
        size -= 0.5
    return size


# ── building ─────────────────────────────────────────────────────────────────

def _deck_builder():
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    class Builder:
        def __init__(self, footer: str):
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
            self.blank = self.prs.slide_layouts[6]
            self.footer = footer
            self.n = 0

        # primitives
        def slide(self, dark: bool = False):
            s = self.prs.slides.add_slide(self.blank)
            self.n += 1
            if dark:
                s.shapes.add_picture(str(BACKGROUND), 0, 0, self.prs.slide_width, self.prs.slide_height)
            else:
                s.background.fill.solid()
                s.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
                self.text(s, self.footer, 0.5, 7.12, 3.0, 0.3, 9, MUTED)
                self.text(s, str(self.n), 12.4, 7.12, 0.5, 0.3, 9, MUTED, align="right")
            return s

        def text(self, s, text, x, y, w, h, pt, color, bold=False, align="left", floor=None, anchor="top",
                 spacing=1.0):
            if floor is not None:
                pt = fit(text, w, h, pt, floor, bold)
            box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            tf = box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[anchor]
            for k, line in enumerate(text.split("\n")):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
                para.line_spacing = spacing
                run = para.add_run()
                run.text = line
                run.font.name, run.font.size, run.font.bold = FONT, Pt(pt), bold
                run.font.color.rgb = RGBColor.from_string(color)
            return box

        def shape(self, s, kind, x, y, w, h, color, line=None, shadow=False):
            shp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor.from_string(color)
            if line:
                shp.line.color.rgb = RGBColor.from_string(line)
                shp.line.width = Pt(1)
            else:
                shp.line.fill.background()
            if shadow:
                from pptx.oxml import parse_xml
                shp._element.spPr.append(parse_xml(SHADOW_XML))
            else:
                shp.shadow.inherit = False
            return shp

        def screenshot(self, s, path: Path | None, x, y, w, h, caption=""):
            """Framed screenshot, cropped (top-anchored) to the frame's shape, like the reference deck."""
            self.shape(s, MSO_SHAPE.RECTANGLE, x - 0.03, y - 0.03, w + 0.06, h + 0.06, WHITE, line=FRAME, shadow=True)
            if path and path.exists():
                pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
                iw, ih = Image.open(path).size
                want = w / h
                if iw / ih > want:            # too wide: trim the sides evenly
                    cut = (1 - want * ih / iw) / 2
                    pic.crop_left = pic.crop_right = cut
                else:                         # too tall: keep the top of the page (header + first content)
                    pic.crop_bottom = 1 - (iw / want) / ih
            if caption:
                self.text(s, caption, x, y + h + 0.13, w, 0.35, 10.5, MUTED, floor=9)

        def logo(self, s, x, y):
            self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 0.62, 0.62, ACCENT)
            tri = self.shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, x + 0.22, y + 0.17, 0.21, 0.19, WHITE)
            tri.rotation = 90

        def head(self, s, kicker, title, w=10.0, lines=1) -> float:
            """Kicker + title; returns the y (inches) where the title's text ends."""
            self.text(s, kicker.upper(), 0.5, 0.5, 6.0, 0.32, 12.5, ACCENT, bold=True)
            h = 0.55 if lines == 1 else 1.05
            size = fit(title, w, h, 30, 22, bold=True)
            self.text(s, title, 0.5, 0.95, w, h, size, INK, bold=True)
            return 0.95 + _lines(title, size, True, w) * size * 1.2 / 72

        # slides
        def title_slide(self, c):
            s = self.slide(dark=True)
            self.logo(s, 6.36, 1.55)
            self.text(s, c["app_name"], 0.5, 2.35, 12.33, 1.1, 54, WHITE, bold=True, align="center", floor=36)
            if c["tagline"]:
                self.text(s, c["tagline"], 0.5, 3.5, 12.33, 0.6, 19, LAVENDER, align="center", floor=15)
            if c["pitch"]:
                self.text(s, c["pitch"], 2.17, 4.35, 9.0, 0.8, 13.5, PALE, align="center", floor=12)

        def problem_slide(self, c, shot):
            p = c["problem"]
            s = self.slide()
            end = self.head(s, "The problem", p["title"])
            if p["intro"]:
                self.text(s, p["intro"], 0.5, max(1.55, end + 0.1), 8.3, 0.75, 13.5, MUTED, floor=11)
            for k, pt in enumerate(p["points"]):
                y = 2.6 + k * 1.38
                self.shape(s, MSO_SHAPE.OVAL, 0.5, y, 0.46, 0.46, ACCENT)
                self.text(s, str(k + 1), 0.5, y, 0.46, 0.46, 15, WHITE, bold=True, align="center", anchor="middle")
                self.text(s, pt["title"], 1.15, y - 0.05, 4.6, 0.4, 15, INK, bold=True, floor=13)
                self.text(s, pt["body"], 1.15, y + 0.33, 4.65, 0.75, 11.5, MUTED, floor=10)
            self.screenshot(s, shot, 6.35, 2.55, 6.5, 3.05, p["caption"])

        def solution_slide(self, c, shot):
            sol = c["solution"]
            s = self.slide()
            end = self.head(s, "The solution", sol["title"], w=6.0, lines=2)
            self.text(s, sol["body"], 0.5, max(2.2, end + 0.25), 5.6, 1.6, 13.5, MUTED, floor=11)
            for k, stage in enumerate(sol["flow"]):
                y = 4.15 + k * 0.44
                self.shape(s, MSO_SHAPE.OVAL, 0.5, y, 0.14, 0.14, ACCENT)
                self.text(s, stage, 0.78, y - 0.1, 5.0, 0.34, 12.5, INK, bold=True, floor=11)
            self.screenshot(s, shot, 6.75, 0.55, 6.1, 3.43)

        def feature_slide(self, f, shot, k, total):
            s = self.slide()
            kicker = f"{f['kicker'] or 'Feature'} · {k} of {total}" if f["kicker"] else f"Feature {k} of {total}"
            if k % 3 == 0:   # every third feature: a wide screenshot, like the reference "How it works" slide
                end = self.head(s, kicker, f["title"])
                self.text(s, f["body"], 0.5, max(1.55, end + 0.1), 12.3, 0.62, 13, MUTED, floor=11)
                self.screenshot(s, shot, 2.27, 2.55, 8.8, 4.2)
            else:            # text left, screenshot right, like the reference "Step n of 6" slides
                end = self.head(s, kicker, f["title"], w=6.2, lines=2)
                self.text(s, f["body"], 0.5, max(2.0, end + 0.25), 5.9, 2.6, 14, MUTED, floor=12)
                self.screenshot(s, shot, 6.9, 2.0, 6.0, 3.38, f["caption"])

        def benefits_slide(self, c):
            b = c["benefits"]
            s = self.slide()
            self.head(s, "Why it matters", b["title"])
            for k, card in enumerate(b["cards"]):
                x = 0.5 + k * 4.25
                self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.1, 3.95, 3.3, TINT)
                self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.35, 2.45, 0.6, 0.6, ACCENT)
                self.text(s, f"{k + 1:02d}", x + 0.35, 2.45, 0.6, 0.6, 14, WHITE, bold=True, align="center",
                          anchor="middle")
                self.text(s, card["title"], x + 0.35, 3.3, 3.25, 0.7, 16, INK, bold=True, floor=13)
                self.text(s, card["body"], x + 0.35, 4.0, 3.25, 1.3, 11.5, MUTED, floor=10)

        def closing_slide(self, c):
            s = self.slide(dark=True)
            self.logo(s, 6.36, 1.5)
            title = c["closing"]["title"]
            size = fit(title, 12.33, 1.6, 38, 28, bold=True)
            self.text(s, title, 0.5, 2.35, 12.33, 1.6, size, WHITE, bold=True, align="center")
            end = 2.35 + _lines(title, size, True, 12.33) * size * 1.2 / 72
            self.text(s, c["closing"]["subtitle"], 0.5, end + 0.3, 12.33, 0.5, 15, LAVENDER, align="center",
                      floor=12)

    return Builder


def build(content: dict, shots: dict[str, Path], out: Path) -> int:
    Builder = _deck_builder()
    b = Builder(content["app_name"])
    b.title_slide(content)
    if content["problem"]["points"]:
        b.problem_slide(content, shots.get(content["problem"]["screenshot"]))
    b.solution_slide(content, shots.get(content["solution"]["screenshot"]))
    feats = content["features"]
    for k, f in enumerate(feats, start=1):
        b.feature_slide(f, shots.get(f["screenshot"]), k, len(feats))
    if content["benefits"]["cards"]:
        b.benefits_slide(content)
    b.closing_slide(content)
    out.parent.mkdir(parents=True, exist_ok=True)
    b.prs.save(out)
    return b.n


def run_deck(demo: Demo, lang: str, out: Path, transcript: str = "", log=print) -> Path:
    """Write the story, grab screenshots from the recording and build deck.pptx; returns its path."""
    narration = "\n".join(s.say.get(lang, "") for s in demo.segments if s.say)
    content = write_content(demo, lang, transcript or narration, log)
    (out.parent / "deck.json").write_text(json.dumps(content, indent=2, ensure_ascii=False), encoding="utf-8")
    with tempfile.TemporaryDirectory() as td:
        shots = screenshots(demo, lang, Path(td), log)
        dedupe(content, shots, [s.id for s in demo.segments], log)
        (out.parent / "deck.json").write_text(json.dumps(content, indent=2, ensure_ascii=False), encoding="utf-8")
        n = build(content, shots, out)
    log(f"Deck ready: {n} slides -> {out}")
    return out
